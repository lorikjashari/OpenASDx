#!/usr/bin/env python3
"""The hackathon slides (#58): four white slides for a non-technical audience, with speaker notes.
They follow the vTPU talk, the team's own TPU design running the same model, and present this
project as the next step on real hardware, without numbers about the vTPU. Every number comes from
the FPGA logs, through the same parsers as report.py.

    llm/tools/.venv/bin/pip install python-pptx      # once, besides report.py's numpy and matplotlib
    llm/tools/.venv/bin/python llm/tools/slides.py
    llm/tools/.venv/bin/python llm/tools/slides.py --deck "slides/Swiss AI Weeks hackathon 2026.pptx"

Writes slides/openasdx.pptx and its charts in slides/img/. With --deck, it also rebuilds our four
slides (FIRST to FIRST + 3) inside a copy of the team's deck, in that deck's title-and-body layout, and
writes slides/hackathon-openasdx.pptx; every other slide is left as it is. In Google Slides: File ->
Import slides. The team then edited the deck by hand: the presented version is
slides/swiss-ai-weeks-hackathon-2026.pptx, where our slides are 10 to 13, and the talk is
slides/TALK.md.
"""
import argparse
import re
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import report as r
import stories_int8 as si

OUT = r.ROOT / "slides"
IMG = OUT / "img"
INK, GREY, ACCENT, CORE = RGBColor(0x22, 0x22, 0x22), RGBColor(0x66, 0x66, 0x66), "#1f77b4", "#d62728"
FONT = "Arial"


def facts():
    g, c, chk = (r.stories_run(f) for f in ("stories-gemmini.uartlog", "stories-cpu.uartlog", "stories260k.uartlog"))
    assert g["passed"] and c["passed"] and chk["passed"] and chk["compared"]
    assert g["tokens"] == c["tokens"] == chk["tokens"]
    pg, pc = r.profile("stories-gemmini.uartlog"), r.profile("stories-cpu.uartlog")
    assert pg["tokens"] == g["tokens"]
    _, tok = si.load()
    dm = r.re.search(r"(\d+) tokens, \d+ GEMMs, \d+ cycles per token on average: ([\d.]+) tokens/s",
                     r.read("demo/stories260k-demo.uartlog"))
    demo = int(dm.group(1)), dm.group(2)  # the recorded run, as the recording shows it
    hz = g["mhz"] * 1e6
    kv = pg["phases"]["quantize KV cache"] + pg["phases"]["copy KV cache"]
    mult = pg["phases"]["gemm: multiply"]
    copies = sum(pg["phases"][k] for k in r.GEMM_PHASES[1:])
    share = lambda x: 100 * x / pg["total"]
    return {
        "prompt": tok.decode(g["tokens"][1:r.header_dims()["PROMPT_LEN"]]),
        "story": tok.decode(g["tokens"][1:]),
        "gemms": chk["gemms"],
        "g_tps": hz / g["mean"], "c_tps": hz / c["mean"], "speedup": c["mean"] / g["mean"],
        "mhz": g["mhz"], "tps_1ghz": 1e9 / g["mean"],
        "mult_x": pc["phases"]["gemm: multiply"] / mult,
        "acc_pct": share(mult), "kv_pct": share(kv), "copies_pct": share(copies),
        "rest_pct": 100 - share(mult) - share(kv) - share(copies),
        "without_kv_x": pg["total"] / (pg["total"] - kv),
        "ngen": len(g["tokens"]) - r.header_dims()["PROMPT_LEN"],
        "demo_tokens": demo[0], "demo_tps": demo[1],
    }


def style():
    plt.rcParams.update({"font.size": 18, "font.family": "sans-serif", "axes.spines.top": False,
                         "axes.spines.right": False, "axes.spines.left": False,
                         "figure.facecolor": "white", "axes.facecolor": "white"})


def speed_chart(f):
    fig, ax = plt.subplots(figsize=(7, 4.2))
    names = ["processor alone", "with the accelerator"]
    vals = [f["c_tps"], f["g_tps"]]
    bars = ax.bar(names, vals, color=[CORE, ACCENT], width=0.6)
    ax.bar_label(bars, labels=[f"{v:.1f}" for v in vals], padding=6, fontsize=26, fontweight="bold")
    ax.set_ylim(0, max(vals) * 1.3)
    ax.set_yticks([])
    ax.set_title("word pieces per second, on the FPGA", fontsize=18, color="#444444")
    fig.tight_layout()
    fig.savefig(IMG / "speed.png", dpi=200)
    plt.close(fig)


def time_chart(f):
    parts = [("accelerator", f["acc_pct"], ACCENT),
             ("copying data for the accelerator", f["copies_pct"], "#9ecae1"),
             ("memory bookkeeping", f["kv_pct"], "#ff7f0e"),
             ("other processor work", f["rest_pct"], "#bbbbbb")]
    fig, ax = plt.subplots(figsize=(11, 2.6))
    left = 0
    for name, pct, color in parts:
        ax.barh([0], [pct], left=left, color=color, edgecolor="white", linewidth=2, height=0.6)
        if pct < 8:  # too narrow for its label: put it above the bar
            ax.text(left + pct / 2, 0.36, f"{pct:.0f}%", ha="center", va="bottom", fontsize=20, color=color)
        else:
            ax.text(left + pct / 2, 0, f"{pct:.0f}%", ha="center", va="center", fontsize=20,
                    color="white" if color in (ACCENT, "#ff7f0e") else "black")
        left += pct
    ax.set_xlim(0, 100)
    ax.set_ylim(-0.5, 0.65)
    ax.axis("off")
    ax.legend([plt.Rectangle((0, 0), 1, 1, color=c) for _, _, c in parts], [n for n, _, _ in parts],
              frameon=False, ncol=4, fontsize=14, loc="upper center", bbox_to_anchor=(0.5, 0.05))
    ax.set_title("where the time goes, for each word piece", fontsize=18, color="#444444")
    fig.tight_layout()
    fig.savefig(IMG / "time.png", dpi=200, bbox_inches="tight")
    plt.close(fig)


def poster_gif(src, dst):
    """The recording, starting on its last frame (the finished story) instead of an empty terminal,
    so the slide's thumbnail and a still export show the story. It loops as before."""
    from PIL import Image, ImageSequence
    im = Image.open(src)
    frames = [fr.copy() for fr in ImageSequence.Iterator(im)]
    durations = [fr.info.get("duration", 100) for fr in ImageSequence.Iterator(Image.open(src))]
    frames, durations = frames[-1:] + frames[:-1], durations[-1:] + durations[:-1]
    frames[0].save(dst, save_all=True, append_images=frames[1:], duration=durations, loop=0,
                   disposal=2, optimize=False)


def text_box(slide, x, y, w, h, lines, size=24, color=INK, bold_first=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.space_after = Pt(14)
        for run in p.runs:
            run.font.size, run.font.name, run.font.color.rgb = Pt(size), FONT, color
            run.font.bold = bold_first and i == 0
    return tf


def table(slide, x, y, w, rows, col_w, size=20):
    """A plain table on white: a bold grey header row, then white rows, dark text."""
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                               Inches(0.62 * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE) if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name = Pt(size), FONT
                    run.font.color.rgb = GREY if j == 0 and i else INK
                    run.font.bold = i == 0
    return t


def slide(prs, title, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    text_box(s, 0.7, 0.45, 12, 1.0, [title], size=38, bold_first=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


# ---------------------------------------------------------------- the team's deck

FIRST = 11  # our slides are 11 to 14 of the team's deck


def bullets(tf, lines, size=None):
    """Dash bullets as the deck writes them; **text** is bold, as the deck marks its key words."""
    from pptx.oxml.ns import qn
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "457200")
        pPr.set("indent", "-342900")
        bu = pPr.makeelement(qn("a:buChar"), {"char": "-"})
        pPr.append(bu)
        p.space_after = Pt(8)
        for k, part in enumerate(re.split(r"\*\*", line)):
            if part:
                run = p.add_run()
                run.text = part
                run.font.bold = k % 2 == 1
                if size:
                    run.font.size = Pt(size)


def deck_slide(prs, title, notes, body=None, box=None, size=None):
    """A slide in the deck's title-and-body layout. box moves the body (inches: x, y, w, h)."""
    layout = next(l for l in prs.slide_layouts if l.name == "TITLE_AND_BODY")
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = title
    ph = s.placeholders[1]
    if body:
        if box:
            ph.left, ph.top, ph.width, ph.height = (Inches(v) for v in box)
        bullets(ph.text_frame, body, size)
    else:
        ph._element.getparent().remove(ph._element)
    s.notes_slide.notes_text_frame.text = notes
    return s


def deck_table(s, x, y, w, rows, col_w, size=13):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                           Inches(0.42 * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2) if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.text = text
            for par in cell.text_frame.paragraphs:
                for run in par.runs:
                    run.font.size, run.font.name = Pt(size), FONT
                    run.font.color.rgb = GREY if j == 0 and i else INK
                    run.font.bold = i == 0


def copy_background(src, dst):
    """The deck gives every slide its own background image; give dst the one src has."""
    import copy
    from pptx.oxml.ns import qn
    bg = src._element.find(qn("p:cSld")).find(qn("p:bg"))
    if bg is None:
        return
    bg = copy.deepcopy(bg)
    for blip in bg.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        blip.set(qn("r:embed"), dst.part.relate_to(src.part.related_part(rid), src.part.rels[rid].reltype))
    csld = dst._element.find(qn("p:cSld"))
    old = csld.find(qn("p:bg"))
    if old is not None:
        csld.remove(old)
    csld.insert(0, bg)


def replace_slides(prs, first, new):
    """Put the new slides (added at the end) in place of slides first.. first+len(new)-1, with
    the backgrounds of the slides they replace."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    old = ids[first - 1:first - 1 + len(new)]
    for k, s in enumerate(new):
        copy_background(prs.slides[first - 1 + k], s)
    for el in old:
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    for k, el in enumerate(ids[-len(new):]):
        lst.remove(el)
        lst.insert(first - 1 + k, el)


def into_deck(f, src, dst):
    prs = Presentation(src)
    new = []
    s = deck_slide(prs, "From our own accelerator to a mature open one", (
        "(About 1 minute.) You just saw the vTPU: an accelerator our team built from scratch, to learn how "
        "a TPU works inside. Then we took the next step with the same model and the same prompt: we took "
        "Gemmini, a more mature open-source accelerator from UC Berkeley that works like a TPU, and ran it "
        "on real hardware: an FPGA, a chip that can be rewired to become any circuit, rented from AWS in "
        "Frankfurt."))
    deck_table(s, 0.45, 1.35, 9.1, [
        ["", "vTPU (the talk before)", "this project"],
        ["the accelerator", "built from scratch by our team", "Gemmini: open source, more mature (UC Berkeley)"],
        ["the model", "Stories260K, a tiny storyteller", "the same model, the same prompt"],
        ["where it runs", "in simulation", "on real hardware: an FPGA in an AWS data center"],
        ["what it shows", "how an accelerator works, inside", "that an open one runs for real"],
    ], col_w=[2.0, 3.0, 4.1])
    new.append(s)

    s = deck_slide(prs, "It runs on real hardware", (
        f"(About 1 minute.) This is the FPGA, in real time. We give it the start of a story, "
        f"\"{f['prompt']}\", and it continues on its own, word by word. The accelerator does the "
        "multiplications; a small processor next to it does the rest. The chip design and all the software "
        "we wrote and ran are open source; only the rented hardware isn't. If asked whether it's right: it "
        "writes exactly the same story as a laptop, and we checked every one of the "
        f"{f['gemms']:,} multiplications against the processor's own answer."),
        body=["A real FPGA in the cloud, writing a story **word by word**",
              "The chip design and all the software: **open source**"],
        box=(6.35, 1.3, 3.35, 3.5), size=16)
    s.shapes.add_picture(str(IMG / "fpga-run.gif"), Inches(0.35), Inches(1.45), width=Inches(5.9))
    new.append(s)

    s = deck_slide(prs, f"The accelerator makes it {f['speedup']:.1f}× faster", (
        "(About 1 minute.) The model writes in word pieces, so we count those. With the accelerator, the "
        f"FPGA writes {f['g_tps']:.1f} word pieces per second, against {f['c_tps']:.1f} on its processor "
        f"alone. The multiplications themselves, the accelerator's job, get {f['mult_x']:.0f}× faster. "
        f"An FPGA is much slower than a real chip ({f['mhz']:.0f} MHz); the same design made as a real "
        f"chip at 1 GHz would write about {f['tps_1ghz']:.0f} word pieces per second, a projection, not a "
        "measurement. If someone noticed the recording's lower number: it writes a longer story "
        f"({f['demo_tokens']} word pieces instead of {f['ngen']}), and each word piece takes a little "
        f"longer than the one before, so its average is {f['demo_tps']} per second."),
        body=[f"**{f['g_tps']:.1f}** word pieces per second, against {f['c_tps']:.1f} without it",
              f"The multiplications: **{f['mult_x']:.0f}×** faster"],
        box=(6.3, 1.5, 3.4, 3.3), size=16)
    s.shapes.add_picture(str(IMG / "speed.png"), Inches(0.35), Inches(1.3), height=Inches(3.6))
    new.append(s)

    s = deck_slide(prs, "A toy example, but it proves the point", (
        "(About 1 minute.) This is a toy: a tiny model, on an FPGA that imitates a real chip. But it "
        "proves the point. We built an accelerator from scratch, to learn how it works. Then we took a "
        "mature open-source one and ran it on real hardware. And it was possible within a hackathon's "
        "time, without compromising on open source anywhere in the chip design or the software."),
        body=["We built an accelerator **from scratch**, then ran a **mature open-source** one on "
              "**real hardware**",
              "**Possible** within the hackathon's time",
              "**No compromise** on open source, in the chip design or the software"],
        box=(0.6, 1.5, 8.8, 3.4), size=22)
    new.append(s)

    replace_slides(prs, FIRST, new)
    register_notes_master(prs)
    prs.save(dst)
    print(f"wrote {dst}: slides {FIRST} to {FIRST + len(new) - 1} rebuilt")

    # The same four slides alone, in the deck's theme, to import into the shared deck.
    alone = Presentation(dst)
    lst = alone.slides._sldIdLst
    for k, el in enumerate(list(lst)):
        if not FIRST - 1 <= k < FIRST - 1 + len(new):
            alone.part.drop_rel(el.rId)
            lst.remove(el)
    out = dst.with_name("openasdx-slides-11-14.pptx")
    alone.save(out)
    print(f"wrote {out}: our {len(alone.slides)} slides alone")


def register_notes_master(prs):
    """python-pptx writes the notes master but doesn't list it in presentation.xml. PowerPoint and
    Google Slides don't mind; Keynote rejects the file as invalid. Add the list, where the schema
    puts it: right after the slide masters' list."""
    from pptx.oxml.ns import qn
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    pres = prs.part._element
    if pres.find(qn("p:notesMasterIdLst")) is not None:
        return
    rid = next(k for k, rel in prs.part.rels.items() if rel.reltype == RT.NOTES_MASTER)
    lst = pres.makeelement(qn("p:notesMasterIdLst"), {})
    lst.append(lst.makeelement(qn("p:notesMasterId"), {qn("r:id"): rid}))
    pres.find(qn("p:sldMasterIdLst")).addnext(lst)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--deck", help="the team's deck, whose slides 11 to 14 are rebuilt")
    args = ap.parse_args()
    f = facts()
    IMG.mkdir(parents=True, exist_ok=True)
    style()
    speed_chart(f)
    time_chart(f)
    poster_gif(r.ROOT / "img/demo/stories260k-fpga.gif", IMG / "fpga-run.gif")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = slide(prs, "Next: the same model on real hardware", (
        "You just saw the vTPU: our team's own TPU design, running a small storyteller model, and showing "
        "how a TPU works from the inside. We took the next step with the same model and the same prompt: "
        "running it on real hardware. For that we used Gemmini, an open-source AI accelerator from UC "
        "Berkeley that works like a TPU, and put it on an FPGA: a chip that can be rewired to become any "
        "circuit, rented from AWS in Frankfurt. The two fit together: designing an accelerator, and "
        "learning what it takes to run one on real hardware."))
    table(s, 0.7, 1.8, 11.9, [
        ["", "vTPU (the talk before)", "this project"],
        ["the model", "Stories260K, a tiny storyteller", "the same model, the same prompt"],
        ["the accelerator", "the vTPU, designed by our team", "Gemmini, open source, from UC Berkeley"],
        ["where it runs", "in simulation, cycle by cycle", "on an FPGA, in an AWS data center"],
        ["what it shows", "how a TPU works, inside", "what it takes to run one on real hardware"],
    ], col_w=[3.0, 4.4, 4.5])

    s = slide(prs, "It runs on a real chip", (
        f"This is a recording of the FPGA, in real time. We gave it the start of a story, \"{f['prompt']}\", "
        "and it continues on its own, word by word. Every multiplication runs on the accelerator; a small "
        "processor next to it does the rest. If asked whether it's right: it writes exactly the same story "
        "as a simulator and as a laptop, and we checked every one of the "
        f"{f['gemms']:,} multiplications against the processor's own answer: all identical."))
    s.shapes.add_picture(str(IMG / "fpga-run.gif"), Inches(1.4), Inches(1.6), width=Inches(10.5))
    text_box(s, 1.4, 6.1, 10.5, 1, ["An FPGA in an AWS data center, writing a story word by word."],
             size=20, color=GREY)

    s = slide(prs, f"The accelerator makes it {f['speedup']:.1f}× faster", (
        "The model writes in word pieces (tokens), so we count those. With the accelerator, the FPGA writes "
        f"{f['g_tps']:.1f} word pieces per second, against {f['c_tps']:.1f} on its processor alone. The "
        f"multiplications themselves, the accelerator's job, get {f['mult_x']:.0f}× faster. The FPGA only "
        f"runs at {f['mhz']:.0f} MHz, far slower than a real chip. The same design made as a real chip at "
        f"1 GHz would take the same steps, and write about {f['tps_1ghz']:.0f} word pieces per second: "
        "a projection, not a measurement. If someone noticed the recording's lower number: it writes a "
        f"longer story ({f['demo_tokens']} word pieces instead of {f['ngen']}), and every word piece takes a "
        f"little longer than the one before, so its average is {f['demo_tps']} per second."))
    s.shapes.add_picture(str(IMG / "speed.png"), Inches(0.5), Inches(1.7), width=Inches(8.0))
    text_box(s, 8.9, 2.6, 4.0, 4, [
        f"{f['g_tps']:.1f} word pieces per second, against {f['c_tps']:.1f}.",
        f"The multiplications: {f['mult_x']:.0f}× faster."], size=24)

    s = slide(prs, "What we learned", (
        f"We measured where the time goes. The accelerator finishes its part in {f['acc_pct']:.0f}% of the "
        "time. The rest is the processor moving and preparing data for it. About half, "
        f"{f['kv_pct']:.0f}%, is re-preparing the model's memory of the story so far, at every word. So the "
        "lesson for every TPU design, ours included: the multiplier is the easy part, and feeding it is the "
        "hard part. Known software fixes could make ours up to "
        f"{f['without_kv_x']:.0f}× faster, and the vTPU's idea of doing more of the steps on the chip itself "
        "points the same way. And every piece we used is open source: the processor, the accelerator and the tools "
        "to put them on the FPGA."))
    s.shapes.add_picture(str(IMG / "time.png"), Inches(0.9), Inches(1.5), width=Inches(11.5))
    text_box(s, 0.9, 4.5, 11.5, 2.8, [
        "The multiplier is the easy part. Feeding it data is the hard part.",
        f"The accelerator works {f['acc_pct']:.0f}% of the time; the rest is the processor's work around it. "
        f"Known fixes: up to {f['without_kv_x']:.0f}× faster.",
        "Everything we used is open source: the processor, the accelerator and the tools."], size=22)

    register_notes_master(prs)
    OUT.mkdir(exist_ok=True)
    prs.save(OUT / "openasdx.pptx")
    print(f"wrote {OUT / 'openasdx.pptx'} and 2 charts in {IMG}")
    if args.deck:
        into_deck(f, args.deck, OUT / "hackathon-openasdx.pptx")


if __name__ == "__main__":
    main()

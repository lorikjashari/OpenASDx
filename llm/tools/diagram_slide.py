#!/usr/bin/env python3
"""A one-slide deck for questions: what we built, end to end, in big blocks (#58). White, 16:9 at the
team deck's size, with editable shapes, so the text can be changed after importing it.

    llm/tools/.venv/bin/python llm/tools/diagram_slide.py        writes slides/openasdx-diagram.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slides import register_notes_master

OUT = Path(__file__).resolve().parents[2] / "slides" / "openasdx-diagram.pptx"
FONT = "Arial"
INK, GREY = RGBColor(0x22, 0x22, 0x22), RGBColor(0x55, 0x55, 0x55)
rgb = lambda h: RGBColor.from_string(h.lstrip("#"))


def box(s, x, y, w, h, title, body=None, fill="#FFFFFF", line="#999999", dash=False, title_size=14,
        body_size=11, anchor_top=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(line)
    shp.line.width = Pt(1.5)
    if dash:
        ln = shp.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = 1 if anchor_top else 3  # top, or middle
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.08))
    tf.margin_top = Inches(0.06)
    lines = [(title, title_size, True, INK)] + ([(body, body_size, False, GREY)] if body else [])
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, FONT, color
    return shp


def arrow(s, x1, y1, x2, y2, label=None, lx=None, ly=None, lw=1.6):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb("#444444")
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if label:
        t = s.shapes.add_textbox(Inches(lx), Inches(ly), Inches(lw), Inches(0.4)).text_frame
        t.word_wrap = True
        p = t.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size, r.font.italic, r.font.name, r.font.color.rgb = Pt(10), True, FONT, GREY


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t = s.shapes.add_textbox(Inches(0.35), Inches(0.2), Inches(9.3), Inches(0.6)).text_frame
    r = t.paragraphs[0].add_run()
    r.text = "OpenASDx, end to end"
    r.font.size, r.font.name, r.font.color.rgb = Pt(26), FONT, INK

    # The laptop: the model, made to fit the accelerator, and the program that runs it.
    box(s, 0.3, 1.0, 2.75, 4.35, "Our laptop", fill="#F3F8F3", line="#6AA36F", dash=True, anchor_top=True)
    box(s, 0.5, 1.5, 2.35, 0.95, "The model", "Stories260K, a tiny storyteller (float)", fill="#FFFFFF", line="#6AA36F")
    box(s, 0.5, 2.8, 2.35, 0.95, "Quantized to 8-bit", "fits the accelerator's integer math (Python)",
        fill="#FFFFFF", line="#6AA36F")
    box(s, 0.5, 4.1, 2.35, 0.95, "The program", "runs the model in C, sends every multiplication to Gemmini",
        fill="#FFFFFF", line="#6AA36F")
    arrow(s, 1.675, 2.45, 1.675, 2.8)
    arrow(s, 1.675, 3.75, 1.675, 4.1)

    # Chipyard: the open-source chip generator the chip and the tools come from.
    box(s, 3.45, 1.0, 3.0, 1.3, "Chipyard (open source)",
        "UC Berkeley's chip generator: the Rocket processor, the Gemmini accelerator, the compiler and FireSim",
        fill="#F5F1FA", line="#8E6BBE")

    # AWS: the manager machine drives the FPGA machine.
    box(s, 3.3, 2.65, 6.4, 2.7, "AWS cloud, Frankfurt", fill="#F6F6F6", line="#999999", dash=True, anchor_top=True)
    box(s, 3.5, 3.15, 2.55, 1.95, "FireSim manager",
        "a cloud computer: builds the program, loads the chip onto the FPGA, starts each run",
        fill="#FFFFFF", line="#777777")
    fpga = box(s, 6.45, 3.15, 3.05, 1.95, "FPGA (AWS F2)", "rewired to become the chip:",
               fill="#EEF4FB", line="#1F77B4", anchor_top=True)
    box(s, 6.6, 4.0, 1.3, 0.9, "Rocket", "RISC-V processor", fill="#FFFFFF", line="#D62728", title_size=12,
        body_size=10)
    box(s, 8.05, 4.0, 1.3, 0.9, "Gemmini", "16×16 accelerator", fill="#FFFFFF", line="#1F77B4", title_size=12,
        body_size=10)

    # The result, back to us.
    box(s, 6.85, 1.0, 2.85, 1.3, "The result", "the story, word by word, and the cycle counts behind the report",
        fill="#FFF8EC", line="#E08A1E")

    arrow(s, 2.85, 4.575, 3.5, 4.3, "the code", 2.75, 4.62, 0.9)
    arrow(s, 4.95, 2.3, 4.8, 3.15, "tools", 4.85, 2.55, 0.8)
    arrow(s, 6.05, 4.1, 6.45, 4.1, "chip + program", 5.6, 3.72, 1.3)
    arrow(s, 7.97, 3.15, 8.1, 2.3, "story, cycles", 8.1, 2.6, 1.3)
    del fpga

    s.notes_slide.notes_text_frame.text = (
        "For questions. On our laptop, we took the model, converted it to 8-bit integers so it fits "
        "Gemmini's math, and wrote the C program that runs it. Chipyard, Berkeley's open-source chip "
        "generator, provides the chip (the Rocket processor and the Gemmini accelerator) and the tools, "
        "including FireSim. In AWS's Frankfurt data center, the FireSim manager builds the program, loads "
        "the chip onto the FPGA and starts each run. The FPGA becomes the chip and writes the story, and "
        "we get the story and the cycle counts back.")
    register_notes_master(prs)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
